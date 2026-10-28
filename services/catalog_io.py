"""Ассортимент в Excel: выгрузка, проверка и загрузка файла админом.

Товар, которого нет в файле, не удаляется, а скрывается (active = 0): на него
могут ссылаться старые заказы, и «История заказов» должна их показывать.
"""
import sqlite3
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "Ассортимент"
HEADERS = ["ID", "Раздел", "Подраздел", "Товар", "Поставщик"]
WIDTHS = [8, 22, 22, 48, 20]
SECTION_DECOR = "🔹"


def _clean(value) -> str:
    # лишние пробелы по краям и внутри — частая причина «дублей»
    return " ".join(str(value if value is not None else "").split())


def _key(value) -> str:
    return _clean(value).replace(SECTION_DECOR, "").strip().lower().replace("ё", "е")


def _plain_section(name: str) -> str:
    return _clean(name.replace(SECTION_DECOR, ""))


def _decor(name: str) -> str:
    # разделы в меню бота оформлены как «🔹 Вино 🔹», в файле — просто «Вино»
    return f"{SECTION_DECOR} {name} {SECTION_DECOR}"


def _root_id(conn: sqlite3.Connection) -> int:
    row = conn.execute(
        "SELECT id FROM categories WHERE COALESCE(parent_id, 0) = 0 ORDER BY id LIMIT 1"
    ).fetchone()
    if row:
        return row[0]
    return conn.execute("INSERT INTO categories (name, parent_id) VALUES ('catalog', 0)").lastrowid


# ---------- выгрузка ----------

def read_catalog(db_path) -> list[dict]:
    """Активные товары в том порядке, в каком их видит пользователь."""
    conn = sqlite3.connect(db_path)
    root = _root_id(conn)
    rows = []

    def items_of(cat_id, section, sub):
        for item_id, name, supplier in conn.execute(
            "SELECT id, name, supplier FROM items WHERE category_id = ? AND active = 1 ORDER BY pos, id",
            (cat_id,)
        ):
            rows.append({"id": item_id, "section": section, "sub": sub,
                         "name": _clean(name), "supplier": _clean(supplier)})

    sections = conn.execute(
        "SELECT id, name FROM categories WHERE parent_id = ? AND active = 1 ORDER BY pos, id", (root,)
    ).fetchall()
    for sec_id, sec_name in sections:
        section = _plain_section(sec_name)
        items_of(sec_id, section, "")
        for sub_id, sub_name in conn.execute(
            "SELECT id, name FROM categories WHERE parent_id = ? AND active = 1 ORDER BY pos, id", (sec_id,)
        ).fetchall():
            items_of(sub_id, section, _clean(sub_name))
    conn.close()
    return rows


def build_workbook(rows: list[dict], out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for r in rows:
        ws.append([r["id"], r["section"], r["sub"], r["name"], r["supplier"]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, width in enumerate(WIDTHS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    # ID серым: его не правят, по нему бот узнаёт товар
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].font = Font(color="808080")
        row[0].alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 2)}"

    help_ws = wb.create_sheet("Как заполнять")
    help_ws.column_dimensions["A"].width = 100
    for line in HELP_LINES:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True, size=13)

    wb.save(out_path)


HELP_LINES = [
    "Как править ассортимент",
    "",
    "Одна строка — один товар. Порядок строк = порядок кнопок в боте.",
    "",
    "• ID — не трогайте. По нему бот понимает, что это тот же товар (и не ломает историю заказов).",
    "• Новый товар — добавьте строку и оставьте ID пустым.",
    "• Убрать товар — удалите строку. Он пропадёт из каталога и поиска, но останется в истории заказов.",
    "• Переименовать или перенести в другой раздел — просто исправьте название / раздел в той же строке.",
    "• Раздел — верхний уровень меню (Вино, Пиво…). Значки 🔹 бот добавит сам.",
    "• Подраздел — второй уровень (Вино белое…). Можно оставить пустым — товар будет прямо в разделе.",
    "• Новый раздел или подраздел появится сам, пустые — пропадут.",
    "",
    "Готовый файл отправьте боту (как документ). Бот покажет, что изменится, и попросит подтвердить.",
    "Перед применением бот сохраняет копию базы.",
]


def export_catalog(db_path, out_path):
    build_workbook(read_catalog(db_path), out_path)


# ---------- загрузка ----------

@dataclass
class ImportPlan:
    rows: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    new_items: list[str] = field(default_factory=list)
    changed: list[str] = field(default_factory=list)
    hidden: list[str] = field(default_factory=list)
    new_categories: list[str] = field(default_factory=list)
    removed_categories: list[str] = field(default_factory=list)

    @property
    def has_changes(self) -> bool:
        return bool(self.new_items or self.changed or self.hidden
                    or self.new_categories or self.removed_categories)


def parse_workbook(path) -> tuple[list[dict], list[str]]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return [], ["Не удалось открыть файл. Нужен Excel-файл .xlsx."]
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # шапку ищем по названиям — столбцы можно переставлять
    header_idx, columns = None, {}
    for i, row in enumerate(all_rows[:5]):
        found = {_key(v): n for n, v in enumerate(row) if _key(v) in {"id", "раздел", "подраздел", "товар", "поставщик"}}
        if "раздел" in found and "товар" in found:
            header_idx, columns = i, found
            break
    if header_idx is None:
        return [], ["Не нашёл шапку таблицы: нужны столбцы «Раздел» и «Товар»."]

    def cell(row, name):
        n = columns.get(name)
        return row[n] if n is not None and n < len(row) else None

    rows, errors = [], []
    seen_ids, seen_names = {}, {}
    for excel_row, row in enumerate(all_rows[header_idx + 1:], start=header_idx + 2):
        raw_id = cell(row, "id")
        section = _plain_section(_clean(cell(row, "раздел")))
        sub = _clean(cell(row, "подраздел"))
        name = _clean(cell(row, "товар"))
        supplier = _clean(cell(row, "поставщик"))
        if not any([_clean(raw_id), section, sub, name, supplier]):
            continue

        item_id = None
        if _clean(raw_id):
            try:
                as_float = float(str(raw_id).replace(",", "."))
                if not as_float.is_integer():
                    raise ValueError
                item_id = int(as_float)
            except ValueError:
                errors.append(f"Строка {excel_row}: ID «{_clean(raw_id)}» — не число. Для нового товара оставьте ID пустым.")
                continue
        if not name:
            errors.append(f"Строка {excel_row}: не заполнено название товара.")
            continue
        if not section:
            errors.append(f"Строка {excel_row}: у «{name}» не заполнен раздел.")
            continue
        if item_id is not None:
            if item_id in seen_ids:
                errors.append(f"Строка {excel_row}: ID {item_id} уже есть в строке {seen_ids[item_id]} (скопированная строка?). У нового товара ID должен быть пустым.")
                continue
            seen_ids[item_id] = excel_row
        name_key = (_key(section), _key(sub), _key(name))
        if name_key in seen_names:
            errors.append(f"Строка {excel_row}: «{name}» уже есть в строке {seen_names[name_key]} в том же разделе.")
            continue
        seen_names[name_key] = excel_row
        rows.append({"row": excel_row, "id": item_id, "section": section, "sub": sub,
                     "name": name, "supplier": supplier})

    if not rows and not errors:
        errors.append("В файле нет ни одного товара.")
    return rows, errors


def _paths(conn, root) -> dict[int, tuple[str, str]]:
    """category_id → (раздел, подраздел) для всех категорий, включая скрытые."""
    cats = {cid: (parent, name) for cid, parent, name in
            conn.execute("SELECT id, COALESCE(parent_id, 0), name FROM categories")}
    out = {}
    for cid, (parent, name) in cats.items():
        if parent == root:
            out[cid] = (_plain_section(name), "")
        elif parent in cats and cats[parent][0] == root:
            out[cid] = (_plain_section(cats[parent][1]), _clean(name))
    return out


def plan_import(db_path, path) -> ImportPlan:
    plan = ImportPlan()
    plan.rows, plan.errors = parse_workbook(path)

    conn = sqlite3.connect(db_path)
    root = _root_id(conn)
    paths = _paths(conn, root)
    cat_active = dict(conn.execute("SELECT id, active FROM categories"))
    existing = {iid: (name, supplier, cat, active) for iid, name, supplier, cat, active in
                conn.execute("SELECT id, name, supplier, category_id, active FROM items")}
    conn.close()

    for r in plan.rows:
        if r["id"] is not None and r["id"] not in existing:
            plan.errors.append(f"Строка {r['row']}: товара с ID {r['id']} нет в базе. Для нового товара оставьте ID пустым.")
    if plan.errors:
        return plan

    # активные разделы сейчас: ключ (раздел, подраздел) → как написано
    active_titles = {}
    for cid, (sec, sub) in paths.items():
        if cat_active.get(cid):
            active_titles[(_key(sec), _key(sub))] = sub or sec

    file_titles, file_ids = {}, set()
    for r in plan.rows:
        where = r["section"] + (f" / {r['sub']}" if r["sub"] else "")
        cat_keys = [((_key(r["section"]), ""), r["section"], r["section"])]
        if r["sub"]:
            cat_keys.append(((_key(r["section"]), _key(r["sub"])), r["sub"], where))
        for k, title, full in cat_keys:
            if k in file_titles:
                continue
            file_titles[k] = title
            if k not in active_titles:
                plan.new_categories.append(full)
            elif active_titles[k] != title:
                plan.changed.append(f"раздел «{active_titles[k]}» → «{title}»")

        if r["id"] is None:
            plan.new_items.append(f"{r['name']} ({where})")
            continue
        file_ids.add(r["id"])
        old_name, old_supplier, old_cat, active = existing[r["id"]]
        old_sec, old_sub = paths.get(old_cat, ("?", ""))
        diffs = []
        if _clean(old_name) != r["name"]:
            diffs.append(f"название: {_clean(old_name)} → {r['name']}")
        if (_key(old_sec), _key(old_sub)) != (_key(r["section"]), _key(r["sub"])):
            diffs.append(f"раздел: {old_sec}{' / ' + old_sub if old_sub else ''} → {where}")
        if _clean(old_supplier) != r["supplier"]:
            diffs.append(f"поставщик: {_clean(old_supplier) or '—'} → {r['supplier'] or '—'}")
        if not active:
            diffs.append("снова в каталоге")
        if diffs:
            plan.changed.append(f"{r['name']}: " + "; ".join(diffs))

    plan.hidden = [_clean(name) for iid, (name, _, _, active) in existing.items()
                   if active and iid not in file_ids]
    for k, title in active_titles.items():
        if k not in file_titles:
            plan.removed_categories.append(title)
    return plan


def apply_import(db_path, rows: list[dict]):
    """Приводит каталог в базе к содержимому файла. Всё в одной транзакции."""
    conn = sqlite3.connect(db_path)
    try:
        root = _root_id(conn)
        by_key = {}
        for cid, parent, name in conn.execute("SELECT id, COALESCE(parent_id, 0), name FROM categories"):
            by_key.setdefault((parent, _key(name)), cid)

        conn.execute("UPDATE categories SET active = 0 WHERE id != ?", (root,))
        conn.execute("UPDATE items SET active = 0")

        positions = {}

        def category(parent, title, display):
            pos = positions.setdefault(parent, {}).setdefault(_key(title), len(positions[parent]))
            cid = by_key.get((parent, _key(title)))
            if cid is None:
                cid = conn.execute(
                    "INSERT INTO categories (name, parent_id, active, pos) VALUES (?, ?, 1, ?)",
                    (display, parent, pos)
                ).lastrowid
                by_key[(parent, _key(title))] = cid
            else:
                conn.execute("UPDATE categories SET name = ?, active = 1, pos = ? WHERE id = ?",
                             (display, pos, cid))
            return cid

        for pos, r in enumerate(rows):
            target = category(root, r["section"], _decor(r["section"]))
            if r["sub"]:
                target = category(target, r["sub"], r["sub"])
            if r["id"] is None:
                conn.execute(
                    "INSERT INTO items (name, supplier, category_id, active, pos) VALUES (?, ?, ?, 1, ?)",
                    (r["name"], r["supplier"] or None, target, pos)
                )
            else:
                conn.execute(
                    "UPDATE items SET name = ?, supplier = ?, category_id = ?, active = 1, pos = ? WHERE id = ?",
                    (r["name"], r["supplier"] or None, target, pos, r["id"])
                )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
